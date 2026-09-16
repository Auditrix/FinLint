"""Read data from an Excel workbook."""

from pathlib import Path

from openpyxl import load_workbook

from finlint.parsing.model import Cell, Sheet, Workbook


def get_formula_details(cell):
    """Get the formula and check whether it is an array formula."""
    if cell.data_type != "f":
        return None, None

    # Most formulas come from openpyxl as normal strings.
    if isinstance(cell.value, str):
        return cell.value, None

    # Array formulas come as objects, so I read their text and range separately.
    formula = cell.value.text
    array_range = cell.value.ref
    return formula, array_range


def get_defined_names(excel_workbook):
    """Get workbook names such as Revenue and the cell they refer to."""
    defined_names = {}

    for defined_name in excel_workbook.defined_names.values():
        # A local name belongs to only one sheet, but our model stores global names.
        if defined_name.localSheetId is not None:
            continue

        if defined_name.attr_text is not None:
            defined_names[defined_name.name] = defined_name.attr_text

    return defined_names


def read_workbook(path: str | Path) -> Workbook:
    """Read an Excel file and return it using our shared models."""
    file_path = Path(path)

    # Checking the path here gives a clearer message than openpyxl would give.
    if not file_path.exists():
        raise FileNotFoundError(f"Workbook not found: {file_path}")

    if file_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("reader.py supports only .xlsx and .xlsm files")

    keep_vba = file_path.suffix.lower() == ".xlsm"

    # I open the same file twice because one version cannot give both values.
    formula_book = load_workbook(
        file_path,
        data_only=False,
        keep_vba=keep_vba,
        keep_links=True,
    )
    value_book = load_workbook(
        file_path,
        data_only=True,
        keep_vba=keep_vba,
        keep_links=True,
    )

    try:
        sheets = {}

        for formula_sheet in formula_book.worksheets:
            value_sheet = value_book[formula_sheet.title]
            cells = {}

            for row in formula_sheet.iter_rows():
                for formula_cell in row:
                    # Blank cells are not useful and large sheets contain many of them.
                    if formula_cell.value is None:
                        continue

                    value_cell = value_sheet[formula_cell.coordinate]
                    formula, array_range = get_formula_details(formula_cell)

                    cell = Cell(
                        sheet=formula_sheet.title,
                        coordinate=formula_cell.coordinate,
                        formula=formula,
                        data_type=formula_cell.data_type,
                        cached_value=value_cell.value,
                        number_format=formula_cell.number_format,
                        array_range=array_range,
                    )
                    cells[formula_cell.coordinate] = cell

            merged_ranges = []
            for cell_range in formula_sheet.merged_cells.ranges:
                merged_ranges.append(str(cell_range))

            sheet = Sheet(
                name=formula_sheet.title,
                state=formula_sheet.sheet_state,
                merged_ranges=merged_ranges,
                cells=cells,
            )
            sheets[formula_sheet.title] = sheet

        return Workbook(
            path=str(file_path),
            sheets=sheets,
            defined_names=get_defined_names(formula_book),
        )
    finally:
        # Closing both files is important even when an error happens above.
        formula_book.close()
        value_book.close()
