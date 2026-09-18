from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import DataTableFormula

from finlint.parsing.reader import read_workbook


def make_workbook(path):
    workbook = OpenpyxlWorkbook()
    sheet = workbook.active
    sheet.title = "Profit and Loss"

    sheet["A1"] = 100
    sheet["A1"].number_format = "#,##0"
    sheet["A2"] = 30
    sheet["A3"] = "=A1-A2"
    sheet.merge_cells("C1:D1")

    hidden_sheet = workbook.create_sheet("Config")
    hidden_sheet.sheet_state = "hidden"
    hidden_sheet["B4"] = 1.1

    revenue_name = DefinedName("Revenue", attr_text="'Profit and Loss'!$A$1")
    workbook.defined_names.add(revenue_name)
    workbook.save(path)


def test_reader_builds_the_shared_models(tmp_path):
    path = tmp_path / "sample.xlsx"
    make_workbook(path)

    result = read_workbook(path)

    profit_sheet = result.sheets["Profit and Loss"]
    formula_cell = profit_sheet.cells["A3"]

    assert formula_cell.formula == "=A1-A2"
    assert formula_cell.data_type == "f"
    assert formula_cell.cached_value is None
    assert formula_cell.calculated_value is None
    assert formula_cell.refs == []
    assert profit_sheet.cells["A1"].cached_value == 100
    assert profit_sheet.cells["A1"].number_format == "#,##0"
    assert profit_sheet.merged_ranges == ["C1:D1"]
    assert result.sheets["Config"].state == "hidden"
    assert result.defined_names["Revenue"] == "'Profit and Loss'!$A$1"


def test_reader_rejects_a_missing_file(tmp_path):
    missing_file = tmp_path / "missing.xlsx"

    try:
        read_workbook(missing_file)
    except FileNotFoundError as error:
        assert "Workbook not found" in str(error)
    else:
        raise AssertionError("A missing workbook should raise FileNotFoundError")


def test_reader_handles_a_data_table_formula(tmp_path):
    path = tmp_path / "data_table.xlsx"
    workbook = OpenpyxlWorkbook()
    sheet = workbook.active
    sheet.title = "WACC and Growth"
    sheet["P9"] = DataTableFormula(ref="P9:Q12", dt2D=True, r1="$B$1", r2="$B$2")
    workbook.save(path)

    result = read_workbook(path)
    cell = result.sheets["WACC and Growth"].cells["P9"]

    assert cell.formula == "=TABLE()"
    assert cell.array_range == "P9:Q12"
