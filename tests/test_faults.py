import json
import random

from openpyxl import Workbook as OpenpyxlWorkbook, load_workbook

from finlint.eval.faults import (
    formula_cells,
    generate_faults,
    load_faults,
    mutate_formula,
    save_fault,
    shift_reference,
    swap_operator,
    write_mutated_copy,
)


def make_workbook(path):
    workbook = OpenpyxlWorkbook()
    sheet = workbook.active
    sheet.title = "Calc"
    sheet["A1"] = 10
    sheet["A2"] = 20
    sheet["A3"] = "=A1+A2"
    sheet["A4"] = "=A3*2"
    sheet["A5"] = "=A4-1"
    workbook.save(path)


def test_swap_operator_changes_one_operator():
    changed = swap_operator("=A1+A2", random.Random(1))
    assert changed == "=A1-A2"


def test_shift_reference_changes_one_row():
    changed = shift_reference("=SUM($A$1)+B2", random.Random(2))
    assert changed != "=SUM($A$1)+B2"
    assert changed.startswith("=")


def test_blank_formula_returns_none():
    assert mutate_formula("=A1+A2", "blank_formula", random.Random(1)) is None


def test_formula_cells_lists_only_formulas(tmp_path):
    path = tmp_path / "sample.xlsx"
    make_workbook(path)
    assert formula_cells(path) == [
        ("Calc", "A3", "=A1+A2"),
        ("Calc", "A4", "=A3*2"),
        ("Calc", "A5", "=A4-1"),
    ]


def test_mutated_copy_does_not_change_source(tmp_path):
    source = tmp_path / "source.xlsx"
    changed = tmp_path / "changed.xlsx"
    make_workbook(source)
    original_bytes = source.read_bytes()

    write_mutated_copy(source, changed, "Calc", "A3", "=A1-A2")

    assert source.read_bytes() == original_bytes
    workbook = load_workbook(changed, data_only=False)
    assert workbook["Calc"]["A3"].value == "=A1-A2"
    workbook.close()


def test_fault_records_are_saved_as_json_lines(tmp_path):
    path = tmp_path / "faults.jsonl"
    record = {"fault_id": 1, "labelled_cell": "Calc!A3"}

    save_fault(path, record)

    assert load_faults(path) == [record]
    assert json.loads(path.read_text(encoding="utf-8")) == record


def test_generate_faults_labels_only_the_injected_cell(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "faults.jsonl"
    make_workbook(source)

    records = generate_faults(source, output, target_faults=2, seed=4, max_attempts=20)

    assert len(records) == 2
    assert all(record["labelled_cell"] == f"{record['sheet']}!{record['coordinate']}" for record in records)
    assert all("changed_cells" not in record for record in records)
