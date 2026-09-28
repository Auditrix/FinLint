"""Build the five test workbooks, each carrying one deliberate problem.

Run from the project root:  python tests/fixtures/make_fixtures.py

The files are small and are committed, so tests never depend on a download.
Every expected finding is written in README.md next to them.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

HERE = Path(__file__).parent


def save(workbook, name):
    path = HERE / name
    workbook.save(path)
    print("wrote %s" % path.name)


def circular():
    """Two cells that read each other, plus a three step loop."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Loop"
    ws["A1"] = "=A2+1"          # A1 and A2 read each other
    ws["A2"] = "=A1+1"
    ws["C1"] = "=C2"            # a longer loop, C1 -> C2 -> C3 -> C1
    ws["C2"] = "=C3"
    ws["C3"] = "=C1"
    ws["E1"] = 10               # a clean cell, must not be reported
    ws["E2"] = "=E1*2"
    save(wb, "circular.xlsx")


def hardcode():
    """A typed number sitting in a row that is otherwise formulas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    for col, value in zip("BCD", (100, 120, 140)):
        ws["%s12" % col] = value            # revenue
        ws["%s13" % col] = value * 0.4      # costs
    ws["B15"] = "=B12-B13"
    ws["C15"] = 72                          # the fault: typed, not calculated
    ws["D15"] = "=D12-D13"
    save(wb, "hardcode.xlsx")


def brokenlink():
    """References that cannot be resolved, plus a stored error value."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "=Missing!B2"                    # sheet does not exist
    ws["A2"] = "=[NotHere.xlsx]Sheet1!A1"       # workbook does not exist
    ws["A3"] = "#REF!"                          # an error left in the file
    ws["A4"] = 5                                # clean
    save(wb, "brokenlink.xlsx")


def arrayformula():
    """One array formula covering a range, next to ordinary formulas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Calc"
    for row in range(1, 4):
        ws["A%d" % row] = row
        ws["B%d" % row] = row * 2
    ws["D1"] = ArrayFormula("D1:D3", "=A1:A3*B1:B3")  # anchor carries the range
    ws["F1"] = "=SUM(A1:A3)"
    save(wb, "arrayformula.xlsx")


def merged():
    """A merged block, where only the top left cell holds anything."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Quarterly summary"
    ws.merge_cells("A1:D1")
    ws["A3"] = 10
    ws["B3"] = 20
    ws["C3"] = "=A3+B3"
    ws.merge_cells("A5:B6")      # a second merge, this one empty
    save(wb, "merged.xlsx")


if __name__ == "__main__":
    circular()
    hardcode()
    brokenlink()
    arrayformula()
    merged()
